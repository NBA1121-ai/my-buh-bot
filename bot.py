"""
Telegram-бот для my-buh
Читает db.json из GitHub (ветка data) и отвечает на вопросы
о складе, финансах, контрагентах, сотрудниках и т.д.
"""

import io
import json
import logging
import os
import re
import time
from datetime import datetime, timedelta
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import requests
from dotenv import load_dotenv
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, BotCommand
from telegram.ext import (
    Application, CommandHandler, MessageHandler, CallbackQueryHandler,
    filters, ContextTypes,
)
from telegram.constants import ParseMode, ChatAction

load_dotenv()

TELEGRAM_TOKEN = os.getenv("TELEGRAM_TOKEN", "")
GITHUB_REPO = os.getenv("GITHUB_REPO", "")
DATA_BRANCH = os.getenv("DATA_BRANCH", "data")
DATA_FILE = os.getenv("DATA_FILE", "db.json")
GITHUB_TOKEN = os.getenv("GITHUB_TOKEN", "")
ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY", "")
NOTIFY_CHAT_IDS = [int(x) for x in os.getenv("NOTIFY_CHAT_IDS", "").split(",") if x.strip()]
NOTIFY_MIN_BALANCE = float(os.getenv("NOTIFY_MIN_BALANCE", "5000"))

try:
    import anthropic
    HAS_ANTHROPIC = True
except ImportError:
    HAS_ANTHROPIC = False

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler()],
)
log = logging.getLogger(__name__)

plt.rcParams["font.family"] = "DejaVu Sans"


# ══════════════════════════════════════════════
#  Загрузка данных с кэшированием
# ══════════════════════════════════════════════

_db_cache = None
_db_cache_time = 0
DB_CACHE_TTL = 60

def fetch_db() -> dict:
    global _db_cache, _db_cache_time
    now = time.time()
    if _db_cache and (now - _db_cache_time) < DB_CACHE_TTL:
        return _db_cache
    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{DATA_FILE}?ref={DATA_BRANCH}"
    headers = {"Accept": "application/vnd.github.v3.raw"}
    if GITHUB_TOKEN:
        headers["Authorization"] = f"token {GITHUB_TOKEN}"
    resp = requests.get(url, headers=headers, timeout=15)
    resp.raise_for_status()
    _db_cache = resp.json()
    _db_cache_time = now
    return _db_cache


# ══════════════════════════════════════════════
#  Стемминг, поиск, форматирование
# ══════════════════════════════════════════════

def stem(word: str) -> str:
    if len(word) <= 3:
        return word
    for suf in ["ами","ями","ов","ев","ей","ах","ях","ом","ем",
                "ой","ей","ам","ям","ий","ый","ая","яя","ое",
                "ее","ую","юю","ые","ие","ок","ек","ик",
                "а","я","о","е","у","ю","ы","и"]:
        if word.endswith(suf) and len(word) - len(suf) >= 3:
            return word[:-len(suf)]
    return word


def fuzzy_match(query: str, name: str) -> bool:
    q_words = [stem(w) for w in query.lower().split() if len(w) > 1]
    if not q_words:
        return False
    name_lower = name.lower()
    name_stems = [stem(w) for w in name_lower.split()]
    for qw in q_words:
        if not any(qw in ns or ns in qw or qw in name_lower for ns in name_stems):
            return False
    return True


def fmt(n: float) -> str:
    if n == int(n):
        return f"{int(n):,}".replace(",", " ")
    return f"{n:,.2f}".replace(",", " ")


_STOP_WORDS = {
    "на","у","нас","в","по","мне","нам","ещё","еще","мы","я","он","они",
    "есть","ли","же","бы","не","и","а","то","от","до","за","из","об",
    "что","как","где","кто","чей","это","вот","тот","так","там",
    "всё","все","весь","общий","общие","итого","только","уже",
    "скажи","покажи","подскажи","расскажи","напиши","дай","выведи",
    "пожалуйста","можно","нужно","хочу","знать","узнать",
    "какой","какая","какие","каков","сколько","много",
}

def clean_search(text: str, extra_stops: list = None) -> str:
    text = re.sub(r'[?!.,;:\-—–()\"\'\«\»]', ' ', text.lower())
    stops = set(_STOP_WORDS)
    if extra_stops:
        for w in extra_stops:
            stops.add(w)
            stops.add(stem(w))
    return " ".join(w for w in text.split() if w not in stops and stem(w) not in stops and len(w) > 1)


# ══════════════════════════════════════════════
#  Справочники → имена по ID
# ══════════════════════════════════════════════

def _lookup(items, item_id, field="name"):
    for item in items:
        if item.get("id") == item_id:
            return item.get(field, item_id)
    return item_id

def get_account_name(db, acc_id):
    raw = acc_id.split(":", 1)[-1] if ":" in acc_id else acc_id
    prefix = acc_id.split(":", 1)[0] if ":" in acc_id else ""
    if prefix == "cash":
        name = _lookup(db.get("cashs", []), raw)
        if name != raw:
            return name
    name = _lookup(db.get("accounts", []), raw)
    return name if name != raw else (raw or "Без счёта")

def get_warehouse_name(db, wh_id):
    return _lookup(db.get("trade", {}).get("warehouses", []), wh_id) if wh_id else "Без склада"

def get_contractor_name(db, c_id):
    return _lookup(db.get("trade", {}).get("contractors", []), c_id)

def get_nom_name(db, n_id):
    return _lookup(db.get("trade", {}).get("nomenclature", []), n_id)

def get_position_name(db, pos_id):
    return _lookup(db.get("trade", {}).get("positions", []), pos_id)


# ══════════════════════════════════════════════
#  Вычисления
# ══════════════════════════════════════════════

_DOC_IN  = {"prihod","postupleniye","receipt","purchase","vozvrat_pokup"}
_DOC_OUT = {"rashod","realizaciya","sale","shipment","vozvrat_post","spisaniye"}

def calc_stock(db):
    stock = {}
    for doc in db.get("trade", {}).get("docs", []):
        dtype = doc.get("type", "")
        is_in, is_out = dtype in _DOC_IN, dtype in _DOC_OUT
        if not is_in and not is_out:
            continue
        wh = doc.get("warehouse", doc.get("warehouseFrom", ""))
        for row in doc.get("rows", []):
            nid = row.get("nomenclature", row.get("nom", ""))
            qty = float(row.get("qty", row.get("quantity", 0)))
            if is_out:
                qty = -qty
            if nid:
                stock.setdefault(nid, {})
                stock[nid][wh] = stock[nid].get(wh, 0) + qty
    return stock


def calc_balances(db):
    bal = {}
    for doc in db.get("bankDocuments", []):
        key = "bank:" + doc.get("account", "")
        s = float(doc.get("sum", 0))
        if doc["type"] == "payment_in":    bal[key] = bal.get(key, 0) + s
        elif doc["type"] == "payment_out": bal[key] = bal.get(key, 0) - s
    for doc in db.get("cashDocuments", []):
        key = "cash:" + doc.get("cash", doc.get("account", ""))
        s = float(doc.get("sum", 0))
        if doc["type"] in ("cash_in", "pko"):    bal[key] = bal.get(key, 0) + s
        elif doc["type"] in ("cash_out", "rko"): bal[key] = bal.get(key, 0) - s
    return bal


def calc_debts(db):
    debts = {}
    for doc in db.get("trade", {}).get("docs", []):
        c_id = doc.get("contractor", "")
        if not c_id:
            continue
        total = sum(float(r.get("total", r.get("sum", 0))) for r in doc.get("rows", []))
        if not total:
            total = float(doc.get("total", doc.get("sum", 0)))
        dtype = doc.get("type", "")
        if dtype in ("realizaciya", "sale", "shipment"):
            debts[c_id] = debts.get(c_id, 0) + total
        elif dtype in ("postupleniye", "prihod", "purchase", "receipt"):
            debts[c_id] = debts.get(c_id, 0) - total
    for doc in db.get("bankDocuments", []) + db.get("cashDocuments", []):
        c_id = doc.get("contractor", "")
        if not c_id:
            continue
        s = float(doc.get("sum", 0))
        dtype = doc.get("type", "")
        if dtype in ("payment_in", "cash_in", "pko"):
            debts[c_id] = debts.get(c_id, 0) - s
        elif dtype in ("payment_out", "cash_out", "rko"):
            debts[c_id] = debts.get(c_id, 0) + s
    return debts


# ══════════════════════════════════════════════
#  Парсинг периодов
# ══════════════════════════════════════════════

_MONTHS_RU = {
    "январ": 1, "феврал": 2, "март": 3, "апрел": 4,
    "ма": 5, "май": 5, "июн": 6, "июл": 7, "август": 8,
    "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12,
}

def parse_period(text: str):
    """Возвращает (date_from, date_to) или None."""
    today = datetime.now().date()
    t = text.lower()

    if "сегодня" in t or "за день" in t:
        return today, today
    if "вчера" in t:
        d = today - timedelta(days=1)
        return d, d
    if "недел" in t:
        start = today - timedelta(days=today.weekday())
        if "прошл" in t:
            start -= timedelta(weeks=1)
            return start, start + timedelta(days=6)
        return start, today
    if "месяц" in t and ("этот" in t or "текущ" in t or not "прошл" in t):
        start = today.replace(day=1)
        return start, today
    if "прошл" in t and "месяц" in t:
        first = today.replace(day=1)
        last_month_end = first - timedelta(days=1)
        last_month_start = last_month_end.replace(day=1)
        return last_month_start, last_month_end

    for key, month_num in _MONTHS_RU.items():
        if key in t:
            year = today.year
            m = re.search(r'(\d{4})', t)
            if m:
                year = int(m.group(1))
            start = datetime(year, month_num, 1).date()
            if month_num == 12:
                end = datetime(year + 1, 1, 1).date() - timedelta(days=1)
            else:
                end = datetime(year, month_num + 1, 1).date() - timedelta(days=1)
            return start, end

    m = re.search(r'(\d{1,2})[./](\d{1,2})[./](\d{2,4})', t)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return datetime(y, mo, d).date(), datetime(y, mo, d).date()
        except ValueError:
            pass

    return None


def filter_docs_by_period(docs, date_from, date_to):
    result = []
    for doc in docs:
        d = doc.get("date", "")
        if not d:
            continue
        try:
            doc_date = datetime.strptime(d[:10], "%Y-%m-%d").date()
        except ValueError:
            continue
        if date_from <= doc_date <= date_to:
            result.append(doc)
    return result


# ══════════════════════════════════════════════
#  Поиск
# ══════════════════════════════════════════════

def find_nomenclature(db, q):
    return [n for n in db.get("trade",{}).get("nomenclature",[]) if fuzzy_match(q, n.get("name",""))]

def find_contractor(db, q):
    return [c for c in db.get("trade",{}).get("contractors",[]) if fuzzy_match(q, c.get("name","")+" "+c.get("full",""))]

def find_employee(db, q):
    return [e for e in db.get("trade",{}).get("employees",[]) if fuzzy_match(q, e.get("name",""))]


# ══════════════════════════════════════════════
#  Форматирование ответов (Telegram Markdown)
# ══════════════════════════════════════════════

def _esc(text: str) -> str:
    return re.sub(r'([_*\[\]()~`>#+\-=|{}.!\\])', r'\\\1', str(text))

def _bold(text: str) -> str:
    return f"*{_esc(text)}*"

def _line(label: str, value: str) -> str:
    return f"  {_esc(label)}: {_bold(value)}"


# ══════════════════════════════════════════════
#  Графики (matplotlib)
# ══════════════════════════════════════════════

def _chart_to_bytes(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="#1e1e2e")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def chart_balances(db) -> bytes | None:
    bal = calc_balances(db)
    if not bal:
        return None
    names = [get_account_name(db, k) for k in bal]
    values = list(bal.values())
    colors = ["#4CAF50" if v >= 0 else "#F44336" for v in values]

    fig, ax = plt.subplots(figsize=(8, max(3, len(names) * 0.8)))
    fig.patch.set_facecolor("#1e1e2e")
    ax.set_facecolor("#1e1e2e")
    bars = ax.barh(names, values, color=colors, edgecolor="#333", height=0.6)
    for bar, val in zip(bars, values):
        ax.text(bar.get_width() + max(values) * 0.02, bar.get_y() + bar.get_height() / 2,
                f"{fmt(val)} сом", va="center", fontsize=10, color="white")
    ax.set_title("Остатки по счетам", fontsize=14, fontweight="bold", color="white", pad=15)
    ax.tick_params(colors="white")
    ax.xaxis.set_major_formatter(ticker.FuncFormatter(lambda x, _: fmt(x)))
    for spine in ax.spines.values():
        spine.set_color("#444")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    return _chart_to_bytes(fig)


def chart_stock(db) -> bytes | None:
    stock = calc_stock(db)
    if not stock:
        return None
    noms = {n["id"]: n for n in db.get("trade", {}).get("nomenclature", [])}
    items = []
    for nid, wh_data in stock.items():
        total = sum(wh_data.values())
        if total > 0:
            n = noms.get(nid, {})
            items.append((n.get("name", nid), total, n.get("unit", "шт")))
    if not items:
        return None
    items.sort(key=lambda x: -x[1])

    names = [f"{i[0]} ({i[2]})" for i in items]
    values = [i[1] for i in items]
    colors = plt.cm.Set3([i / max(len(items), 1) for i in range(len(items))])

    fig, ax = plt.subplots(figsize=(8, max(3, len(items) * 0.8)))
    fig.patch.set_facecolor("#1e1e2e")
    ax.set_facecolor("#1e1e2e")
    bars = ax.barh(names, values, color=colors, edgecolor="#333", height=0.6)
    for bar, val in zip(bars, values):
        ax.text(bar.get_width() + max(values) * 0.02, bar.get_y() + bar.get_height() / 2,
                fmt(val), va="center", fontsize=10, color="white")
    ax.set_title("Остатки на складе", fontsize=14, fontweight="bold", color="white", pad=15)
    ax.tick_params(colors="white")
    for spine in ax.spines.values():
        spine.set_color("#444")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    return _chart_to_bytes(fig)


def chart_debts(db) -> bytes | None:
    debts = calc_debts(db)
    active = {c: d for c, d in debts.items() if d != 0}
    if not active:
        return None
    names = [get_contractor_name(db, c) for c in active]
    values = list(active.values())
    colors = ["#F44336" if v > 0 else "#FF9800" for v in values]
    labels = [f"{fmt(abs(v))} сом\n({'нам должны' if v > 0 else 'мы должны'})" for v in values]

    fig, ax = plt.subplots(figsize=(7, 7))
    fig.patch.set_facecolor("#1e1e2e")
    ax.set_facecolor("#1e1e2e")
    abs_values = [abs(v) for v in values]
    wedges, texts = ax.pie(abs_values, labels=None, colors=colors, startangle=90,
                           wedgeprops={"edgecolor": "#1e1e2e", "linewidth": 2})
    legend_labels = [f"{n}: {l}" for n, l in zip(names, labels)]
    ax.legend(wedges, legend_labels, loc="lower center", bbox_to_anchor=(0.5, -0.15),
              fontsize=9, facecolor="#2a2a3e", edgecolor="#444", labelcolor="white")
    ax.set_title("Взаиморасчёты", fontsize=14, fontweight="bold", color="white", pad=15)
    return _chart_to_bytes(fig)


def chart_income_expense(db, date_from=None, date_to=None) -> bytes | None:
    bank_docs = db.get("bankDocuments", [])
    cash_docs = db.get("cashDocuments", [])
    all_docs = bank_docs + cash_docs
    if date_from and date_to:
        all_docs = filter_docs_by_period(all_docs, date_from, date_to)
    if not all_docs:
        return None

    income, expense = 0, 0
    for doc in all_docs:
        s = float(doc.get("sum", 0))
        dtype = doc.get("type", "")
        if dtype in ("payment_in", "cash_in", "pko"):
            income += s
        elif dtype in ("payment_out", "cash_out", "rko"):
            expense += s

    if income == 0 and expense == 0:
        return None

    fig, ax = plt.subplots(figsize=(6, 5))
    fig.patch.set_facecolor("#1e1e2e")
    ax.set_facecolor("#1e1e2e")
    bars = ax.bar(["Приходы", "Расходы"], [income, expense],
                  color=["#4CAF50", "#F44336"], edgecolor="#333", width=0.5)
    for bar, val in zip(bars, [income, expense]):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + max(income, expense) * 0.02,
                f"{fmt(val)} сом", ha="center", fontsize=11, fontweight="bold", color="white")
    title = "Приходы и расходы"
    if date_from and date_to:
        title += f"\n{date_from.strftime('%d.%m.%Y')} — {date_to.strftime('%d.%m.%Y')}"
    ax.set_title(title, fontsize=13, fontweight="bold", color="white", pad=15)
    ax.tick_params(colors="white")
    ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda x, _: fmt(x)))
    for spine in ax.spines.values():
        spine.set_color("#444")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    return _chart_to_bytes(fig)


# ══════════════════════════════════════════════
#  Экспорт в Excel
# ══════════════════════════════════════════════

_HEADER_FILL = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

def _style_header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER

def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
            cell.border = _BORDER
        ws.column_dimensions[col_letter].width = min(max_len + 3, 40)


def export_balances_xlsx(db) -> bytes | None:
    bal = calc_balances(db)
    if not bal:
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Балансы"
    ws.append(["Счёт / Касса", "Тип", "Остаток (сом)"])
    _style_header(ws, 1, 3)
    total = 0
    for acc_id, amount in bal.items():
        name = get_account_name(db, acc_id)
        acc_type = "Банк" if acc_id.startswith("bank:") else "Касса"
        ws.append([name, acc_type, amount])
        total += amount
    ws.append(["ИТОГО", "", total])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_stock_xlsx(db) -> bytes | None:
    stock = calc_stock(db)
    noms = {n["id"]: n for n in db.get("trade", {}).get("nomenclature", [])}
    if not stock:
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Склад"
    ws.append(["Товар", "Артикул", "Ед.", "Остаток", "Цена", "Себест.", "Сумма"])
    _style_header(ws, 1, 7)
    for nid, wh_data in stock.items():
        total_qty = sum(wh_data.values())
        n = noms.get(nid, {})
        price = n.get("price", 0)
        cost = n.get("cost", 0)
        ws.append([n.get("name", nid), n.get("article", ""), n.get("unit", "шт"),
                    total_qty, price, cost, total_qty * cost])
    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_debts_xlsx(db) -> bytes | None:
    debts = calc_debts(db)
    active = {c: d for c, d in debts.items() if d != 0}
    if not active:
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Взаиморасчёты"
    ws.append(["Контрагент", "Нам должны (сом)", "Мы должны (сом)"])
    _style_header(ws, 1, 3)
    total_owe_us, total_we_owe = 0, 0
    for c_id, amount in sorted(active.items(), key=lambda x: -abs(x[1])):
        name = get_contractor_name(db, c_id)
        owe_us = amount if amount > 0 else 0
        we_owe = -amount if amount < 0 else 0
        ws.append([name, owe_us, we_owe])
        total_owe_us += owe_us
        total_we_owe += we_owe
    ws.append(["ИТОГО", total_owe_us, total_we_owe])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_period_xlsx(db, date_from, date_to) -> bytes | None:
    bank_docs = filter_docs_by_period(db.get("bankDocuments", []), date_from, date_to)
    cash_docs = filter_docs_by_period(db.get("cashDocuments", []), date_from, date_to)
    trade_docs = filter_docs_by_period(db.get("trade", {}).get("docs", []), date_from, date_to)
    if not bank_docs and not cash_docs and not trade_docs:
        return None

    wb = openpyxl.Workbook()

    if bank_docs or cash_docs:
        ws = wb.active
        ws.title = "Деньги"
        ws.append(["Дата", "Номер", "Тип", "Сумма", "Счёт/Касса", "Контрагент", "Назначение"])
        _style_header(ws, 1, 7)
        for doc in sorted(bank_docs + cash_docs, key=lambda d: d.get("date", "")):
            dtype = doc.get("type", "")
            direction = "Приход" if "in" in dtype or "pko" in dtype else "Расход"
            acc = doc.get("account", doc.get("cash", ""))
            ws.append([
                doc.get("date", ""), doc.get("number", ""), direction,
                float(doc.get("sum", 0)), get_account_name(db, f"bank:{acc}"),
                get_contractor_name(db, doc.get("contractor", "")),
                doc.get("purpose", doc.get("note", "")),
            ])
        _auto_width(ws)

    if trade_docs:
        ws2 = wb.create_sheet("Торговля")
        ws2.append(["Дата", "Номер", "Тип", "Контрагент", "Товар", "Кол-во", "Цена", "Сумма"])
        _style_header(ws2, 1, 8)
        for doc in sorted(trade_docs, key=lambda d: d.get("date", "")):
            dtype = "Поступление" if doc.get("type") in _DOC_IN else "Реализация"
            c_name = get_contractor_name(db, doc.get("contractor", ""))
            for row in doc.get("rows", []):
                n_name = get_nom_name(db, row.get("nomenclature", row.get("nom", "")))
                ws2.append([
                    doc.get("date", ""), doc.get("number", ""), dtype, c_name,
                    n_name, float(row.get("qty", 0)),
                    float(row.get("price", 0)), float(row.get("total", 0)),
                ])
        _auto_width(ws2)

    if not bank_docs and not cash_docs:
        wb.remove(wb.active)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ══════════════════════════════════════════════
#  Обработчики вопросов
# ══════════════════════════════════════════════

def handle_summary(db, text):
    kw = ["сводк","отчёт","отчет","итог","обзор","дашборд","dashboard","статус"]
    if not any(k in text for k in kw):
        return None

    bal = calc_balances(db)
    total_money = sum(bal.values())
    stock = calc_stock(db)
    stock_items = sum(1 for nid in stock if sum(stock[nid].values()) > 0)
    debts = calc_debts(db)
    owe_us = sum(v for v in debts.values() if v > 0)
    we_owe = sum(-v for v in debts.values() if v < 0)
    contractors = db.get("trade",{}).get("contractors",[])
    employees = db.get("trade",{}).get("employees",[])
    noms = db.get("trade",{}).get("nomenclature",[])

    lines = [f"{'='*28}", _bold("СВОДКА"), f"{'='*28}\n"]
    lines.append(f"💰 {_bold('Финансы')}")
    for acc_id, amount in bal.items():
        lines.append(f"  {_esc(get_account_name(db, acc_id))}: {_bold(fmt(amount) + ' сом')}")
    lines.append(f"  {'─'*20}")
    lines.append(f"  Итого: {_bold(fmt(total_money) + ' сом')}\n")

    lines.append(f"📦 {_bold('Склад')}: {_esc(str(stock_items))} позиций с остатками")
    lines.append(f"📋 {_bold('Номенклатура')}: {_esc(str(len(noms)))} позиций")
    lines.append(f"🤝 {_bold('Контрагенты')}: {_esc(str(len(contractors)))}")
    lines.append(f"👥 {_bold('Сотрудники')}: {_esc(str(len(employees)))}\n")

    if owe_us or we_owe:
        lines.append(f"📊 {_bold('Взаиморасчёты')}")
        if owe_us:
            lines.append(f"  Нам должны: {_bold(fmt(owe_us) + ' сом')}")
        if we_owe:
            lines.append(f"  Мы должны: {_bold(fmt(we_owe) + ' сом')}")

    return "\n".join(lines)


def handle_money(db, text):
    kw = ["деньг","денег","денежн","баланс","счёт","счет","касс","финанс","бюджет"]
    if "взаиморасч" in text:
        return None
    if not any(k in text for k in kw):
        return None
    bal = calc_balances(db)
    if not bal:
        return _bold("Нет данных по движению денег\\.")

    lines = [f"💰 {_bold('Остатки по счетам')}\n"]
    bank_total, cash_total = 0, 0
    banks, cashes = [], []
    for acc_id, amount in bal.items():
        name = get_account_name(db, acc_id)
        entry = f"  {_esc(name)}: {_bold(fmt(amount) + ' сом')}"
        if acc_id.startswith("bank:"):
            banks.append(entry)
            bank_total += amount
        else:
            cashes.append(entry)
            cash_total += amount
    if banks:
        lines.append(f"🏦 {_esc('Банковские счета')}:")
        lines.extend(banks)
        lines.append("")
    if cashes:
        lines.append(f"💵 {_esc('Кассы')}:")
        lines.extend(cashes)
        lines.append("")
    lines.append(f"{'─'*24}")
    lines.append(f"Итого: {_bold(fmt(bank_total + cash_total) + ' сом')}")
    return "\n".join(lines)


def handle_debts(db, text):
    kw_all = ["должн","должен","долг","долж","дебитор","задолжен","задолженн","взаиморасчёт","взаиморасчет"]
    kw_we = ["мы должн","мы должен","наш долг","кредитор","наша задолженн","кому мы должн","кому должны мы"]
    if not any(k in text for k in kw_all):
        return None

    debts = calc_debts(db)
    if not debts:
        return f"📊 {_bold('Взаиморасчёты')}\n\nДанных нет\\. Долги появятся после оформления документов\\."

    search = clean_search(text, ["должн","долг","дебитор","кредитор","задолженн","контрагент","поставщик","покупател","клиент","нам","мы"])
    if search:
        found = find_contractor(db, search)
        if found:
            lines = [f"📊 {_bold('Взаиморасчёты')}\n"]
            for c in found:
                d = debts.get(c["id"], 0)
                if d > 0:
                    lines.append(f"🔴 {_esc(c['name'])}: должен нам {_bold(fmt(d) + ' сом')}")
                elif d < 0:
                    lines.append(f"🟡 {_esc(c['name'])}: мы должны {_bold(fmt(-d) + ' сом')}")
                else:
                    lines.append(f"🟢 {_esc(c['name'])}: расчёты закрыты")
            return "\n".join(lines)

    owe_us, we_owe = {}, {}
    for c_id, amount in debts.items():
        if amount > 0:   owe_us[c_id] = amount
        elif amount < 0:  we_owe[c_id] = -amount

    we_ask_ours = any(k in text for k in kw_we)
    lines = [f"📊 {_bold('Взаиморасчёты')}\n"]

    if not we_ask_ours and owe_us:
        lines.append(f"🔴 {_esc('Нам должны')}:")
        for c_id, amt in sorted(owe_us.items(), key=lambda x: -x[1]):
            lines.append(f"  {_esc(get_contractor_name(db, c_id))}: {_bold(fmt(amt) + ' сом')}")
        lines.append(f"  Итого: {_bold(fmt(sum(owe_us.values())) + ' сом')}\n")
    if we_owe:
        lines.append(f"🟡 {_esc('Мы должны')}:")
        for c_id, amt in sorted(we_owe.items(), key=lambda x: -x[1]):
            lines.append(f"  {_esc(get_contractor_name(db, c_id))}: {_bold(fmt(amt) + ' сом')}")
        lines.append(f"  Итого: {_bold(fmt(sum(we_owe.values())) + ' сом')}")

    if not owe_us and not we_owe:
        lines.append("🟢 Все расчёты закрыты, долгов нет\\.")
    return "\n".join(lines)


def handle_price(db, text):
    kw = ["стоит","стоимость","цен","прайс","расценк"]
    if not any(k in text for k in kw):
        return None
    search = clean_search(text, ["стоит","стоимость","цена","прайс","расценка","товар"])
    noms = db.get("trade",{}).get("nomenclature",[])
    if search:
        noms = find_nomenclature(db, search)
        if not noms:
            return f"Товар '{_esc(search)}' не найден\\."

    lines = [f"🏷 {_bold('Прайс-лист')}\n"]
    for n in noms:
        p = n.get("price", 0)
        c = n.get("cost", 0)
        if p > 0 or search:
            line = f"  {_esc(n['name'])}: {_bold(fmt(p) + ' сом')}/{_esc(n.get('unit','шт'))}"
            if c > 0 and search:
                line += f" \\(себест\\. {_esc(fmt(c))} сом\\)"
            lines.append(line)
    return "\n".join(lines)


def handle_warehouses(db, text):
    if not re.search(r'\bсклад[ыа]?\b', text):
        return None
    if any(k in text for k in ["сколько","остат","наличи","есть","что на"]):
        return None
    whs = db.get("trade",{}).get("warehouses",[])
    if not whs:
        return "Склады не найдены\\."
    lines = [f"🏭 {_bold('Склады')}\n"]
    for w in whs:
        lines.append(f"  📍 {_bold(w['name'])}")
        lines.append(f"     Тип: {_esc(w.get('kind',''))}")
        lines.append(f"     Ответственный: {_esc(w.get('responsible','-'))}")
    return "\n".join(lines)


def handle_stock(db, text):
    kw = ["склад","остат","наличи","сколько","есть ли","имеется"]
    if not any(k in text for k in kw):
        return None
    search = clean_search(text, [
        "склад","складе","складу","складах","остатки","остаток","остатков",
        "наличие","наличии","имеется","имеются","товар","штук","штуки","кг","шт",
    ])

    noms = db.get("trade",{}).get("nomenclature",[])
    if search:
        found = find_nomenclature(db, search)
        if not found:
            suggestions = [n["name"] for n in noms if any(w in n["name"].lower() for w in search.split() if len(w)>2)]
            if suggestions:
                return f"Товар '{_esc(search)}' не найден\\. Может имелось в виду:\n" + "\n".join(f"  • {_esc(s)}" for s in suggestions[:10])
            return f"Товар '{_esc(search)}' не найден\\."
        noms = found

    stock = calc_stock(db)
    has_any = any(stock.get(n["id"]) for n in noms)

    if not has_any:
        lines = [f"📦 {_bold('Остатки на складе')}\n"]
        show = [n for n in noms if n.get("kind") in ("Товар","Материал","Продукция","Тара")]
        if show:
            for n in show:
                lines.append(f"  {_esc(n['name'])}: {_bold('0')} {_esc(n.get('unit','шт'))}")
            lines.append(f"\n⚠️ {_esc('Движений пока нет. Оформите поступление товара на сайте.')}")
        return "\n".join(lines)

    lines = [f"📦 {_bold('Остатки на складе')}\n"]
    for n in noms:
        ns = stock.get(n["id"], {})
        total = sum(ns.values())
        if total > 0:
            lines.append(f"  ✅ {_esc(n['name'])}: {_bold(fmt(total))} {_esc(n.get('unit','шт'))}")
            if len(ns) > 1:
                for wh_id, qty in ns.items():
                    lines.append(f"       {_esc(get_warehouse_name(db, wh_id))}: {_esc(fmt(qty))}")
        elif total < 0:
            lines.append(f"  ⚠️ {_esc(n['name'])}: {_bold(fmt(total))} {_esc(n.get('unit','шт'))} \\(минус\\!\\)")
        elif search:
            lines.append(f"  ◻️ {_esc(n['name'])}: {_bold('0')} {_esc(n.get('unit','шт'))}")
    return "\n".join(lines)


def handle_goods(db, text):
    kw = ["товар","номенклатур","ассортимент","что продаём","что продаем","каталог"]
    if not any(k in text for k in kw):
        return None
    noms = db.get("trade",{}).get("nomenclature",[])
    if not noms:
        return "Номенклатура пуста\\."

    lines = [f"📋 {_bold('Номенклатура')} \\({_esc(str(len(noms)))} позиций\\)\n"]
    by_kind = {}
    for n in noms:
        by_kind.setdefault(n.get("kind","Прочее"), []).append(n)
    for kind, items in by_kind.items():
        lines.append(f"  {_bold(kind)}:")
        for n in items:
            p = fmt(n.get("price",0))
            lines.append(f"    {_esc(n['name'])} — {_bold(p + ' сом')}/{_esc(n.get('unit','шт'))}")
        lines.append("")
    return "\n".join(lines)


def handle_contractors(db, text):
    kw = ["контрагент","поставщик","покупател","клиент","партнёр","партнер"]
    if not any(k in text for k in kw):
        return None
    search = clean_search(text, kw + ["список","информац","данные","инн"])
    contractors = db.get("trade",{}).get("contractors",[])
    if search:
        contractors = find_contractor(db, search)
    if not contractors:
        return "Контрагенты не найдены\\."

    if len(contractors) <= 3 and search:
        lines = []
        for c in contractors:
            lines.append(f"🤝 {_bold(c.get('name',''))}\n")
            if c.get("full"):    lines.append(_line("Полное название", c["full"]))
            if c.get("inn"):     lines.append(_line("ИНН", c["inn"]))
            if c.get("kind"):    lines.append(_line("Тип", c["kind"]))
            if c.get("address"): lines.append(_line("Адрес", c["address"]))
            if c.get("phone"):   lines.append(_line("Телефон", c["phone"]))
            contracts = [d for d in db.get("trade",{}).get("contracts",[]) if d.get("contractor")==c["id"]]
            if contracts:
                lines.append(f"\n  📝 {_esc('Договоры')}:")
                for d in contracts:
                    lines.append(f"    №{_esc(d.get('number','?'))} от {_esc(d.get('date','?'))} — {_esc(d.get('name',''))}")
        return "\n".join(lines)

    lines = [f"🤝 {_bold('Контрагенты')} \\({_esc(str(len(contractors)))}\\)\n"]
    for c in contractors:
        lines.append(f"  • {_esc(c['name'])}  \\|  ИНН: {_esc(c.get('inn','-'))}")
    return "\n".join(lines)


def handle_employees(db, text):
    kw = ["сотрудник","работник","персонал","кадр","штат","табельн"]
    if not any(k in text for k in kw):
        employees = db.get("trade",{}).get("employees",[])
        search = clean_search(text, ["информация","инфо","данные","расскажи","про","кто","такой","такая"])
        if not (search and any(fuzzy_match(search, e.get("name","")) for e in employees)):
            return None

    search = clean_search(text, kw + ["список","информац","данные"])
    employees = db.get("trade",{}).get("employees",[])
    if search:
        found = find_employee(db, search)
        if found:
            employees = found

    if not employees:
        return "Список сотрудников пуст\\."

    if len(employees) <= 2 and search:
        lines = []
        for e in employees:
            lines.append(f"👤 {_bold(e.get('name',''))}\n")
            lines.append(_line("Таб. №", e.get("tabNo","-")))
            lines.append(_line("Должность", get_position_name(db, e.get("position",""))))
            if e.get("phone"):    lines.append(_line("Телефон", e["phone"]))
            if e.get("birth"):    lines.append(_line("Дата рождения", e["birth"]))
            if e.get("inn"):      lines.append(_line("ИНН", e["inn"]))
            if e.get("address"):  lines.append(_line("Адрес", e["address"]))
            salary = e.get("salary", 0)
            if salary:            lines.append(_line("Оклад", f"{fmt(salary)} сом"))
            if e.get("status"):   lines.append(_line("Статус", e["status"]))
        return "\n".join(lines)

    lines = [f"👥 {_bold('Сотрудники')} \\({_esc(str(len(employees)))}\\)\n"]
    for e in employees:
        pos = get_position_name(db, e.get("position",""))
        lines.append(f"  • {_esc(e['name'])} — {_esc(pos)}")
    return "\n".join(lines)


def handle_org(db, text):
    kw = ["организац","компани","фирм","реквизит","наша компан","юрлиц"]
    if not any(k in text for k in kw):
        return None
    org = db.get("trade",{}).get("org",{})
    if not org or not org.get("name"):
        return "Данные организации не заполнены\\."
    lines = [f"🏢 {_bold('Организация')}\n"]
    if org.get("name"):       lines.append(_line("Название", org["name"]))
    if org.get("inn"):        lines.append(_line("ИНН", org["inn"]))
    if org.get("address"):    lines.append(_line("Адрес", org["address"]))
    if org.get("phone"):      lines.append(_line("Телефон", org["phone"]))
    if org.get("director"):   lines.append(_line("Руководитель", org["director"]))
    if org.get("accountant"): lines.append(_line("Бухгалтер", org["accountant"]))
    return "\n".join(lines)


def handle_payroll(db, text):
    kw = ["зарплат","оклад","налог","отчислен","ставк","соцфонд","подоходн"]
    if not any(k in text for k in kw):
        return None
    pr = db.get("trade",{}).get("payroll",{})
    if not pr:
        return "Зарплатные настройки не заданы\\."
    lines = [f"💼 {_bold('Зарплатные настройки')}\n"]
    lines.append(f"  Подоходный налог: {_bold(str(pr.get('incomeTax',0)) + '%')}")
    lines.append(f"  Соцфонд \\(сотрудник\\): {_bold(str(pr.get('sfEmployee',0)) + '%')}")
    lines.append(f"  Соцфонд \\(работодатель\\): {_bold(str(pr.get('sfEmployer',0)) + '%')}")
    lines.append(f"  Стандартный вычет: {_bold(fmt(pr.get('stdDeduction',0)) + ' сом')}")
    lines.append(f"  Мин\\. зарплата: {_bold(fmt(pr.get('minSalary',0)) + ' сом')}")
    return "\n".join(lines)


def handle_contracts(db, text):
    kw = ["договор","контракт"]
    if not any(k in text for k in kw):
        return None
    contracts = db.get("trade",{}).get("contracts",[])
    if not contracts:
        return "Договоры не найдены\\."
    lines = [f"📝 {_bold('Договоры')} \\({_esc(str(len(contracts)))}\\)\n"]
    for d in contracts:
        c_name = get_contractor_name(db, d.get("contractor",""))
        lines.append(f"  №{_esc(d.get('number','?'))} от {_esc(d.get('date','?'))}")
        lines.append(f"    {_esc(d.get('name',''))} \\| {_esc(c_name)}")
        lines.append("")
    return "\n".join(lines)


def handle_period_report(db, text):
    """Отчёты за период: 'продажи за сентябрь', 'расходы за неделю' и т.д."""
    kw = ["за сегодня","за вчера","за недел","за месяц","за прошл",
          "за январ","за феврал","за март","за апрел","за май","за июн",
          "за июл","за август","за сентябр","за октябр","за ноябр","за декабр",
          "этот месяц","текущий месяц","прошлый месяц","эту неделю","прошлую неделю"]
    if not any(k in text for k in kw):
        return None

    period = parse_period(text)
    if not period:
        return None
    date_from, date_to = period

    bank_docs = filter_docs_by_period(db.get("bankDocuments", []), date_from, date_to)
    cash_docs = filter_docs_by_period(db.get("cashDocuments", []), date_from, date_to)
    trade_docs = filter_docs_by_period(db.get("trade", {}).get("docs", []), date_from, date_to)

    if not bank_docs and not cash_docs and not trade_docs:
        return f"📅 За период {_esc(date_from.strftime('%d.%m.%Y'))} — {_esc(date_to.strftime('%d.%m.%Y'))} операций не найдено\\."

    income = sum(float(d.get("sum", 0)) for d in bank_docs + cash_docs if d.get("type") in ("payment_in", "cash_in", "pko"))
    expense = sum(float(d.get("sum", 0)) for d in bank_docs + cash_docs if d.get("type") in ("payment_out", "cash_out", "rko"))
    trade_in = sum(sum(float(r.get("total", 0)) for r in d.get("rows", [])) for d in trade_docs if d.get("type") in _DOC_IN) or 0
    trade_out = sum(sum(float(r.get("total", 0)) for r in d.get("rows", [])) for d in trade_docs if d.get("type") in _DOC_OUT) or 0

    period_str = f"{date_from.strftime('%d.%m.%Y')} — {date_to.strftime('%d.%m.%Y')}"
    lines = [f"📅 {_bold('Отчёт за период')}", f"  {_esc(period_str)}\n"]

    if income or expense:
        lines.append(f"💰 {_bold('Денежные операции')}")
        lines.append(f"  Приходы: {_bold(fmt(income) + ' сом')} \\({_esc(str(len([d for d in bank_docs + cash_docs if d.get('type') in ('payment_in','cash_in','pko')])))} док\\.\\)")
        lines.append(f"  Расходы: {_bold(fmt(expense) + ' сом')} \\({_esc(str(len([d for d in bank_docs + cash_docs if d.get('type') in ('payment_out','cash_out','rko')])))} док\\.\\)")
        lines.append(f"  Разница: {_bold(fmt(income - expense) + ' сом')}\n")

    if trade_in or trade_out:
        lines.append(f"📦 {_bold('Торговые операции')}")
        lines.append(f"  Поступления: {_bold(fmt(trade_in) + ' сом')}")
        lines.append(f"  Реализации: {_bold(fmt(trade_out) + ' сом')}")
        if trade_out > 0:
            lines.append(f"  Маржа: {_bold(fmt(trade_out - trade_in) + ' сом')}")

    return "\n".join(lines)


# ══════════════════════════════════════════════
#  AI-ассистент (Claude)
# ══════════════════════════════════════════════

def _prepare_data_context(db: dict) -> str:
    parts = []

    bal = calc_balances(db)
    if bal:
        lines = []
        for acc_id, amount in bal.items():
            lines.append(f"  {get_account_name(db, acc_id)}: {fmt(amount)} сом")
        parts.append("БАЛАНСЫ СЧЕТОВ:\n" + "\n".join(lines))

    stock = calc_stock(db)
    noms = {n["id"]: n for n in db.get("trade", {}).get("nomenclature", [])}
    if stock:
        lines = []
        for nid, wh_data in stock.items():
            total = sum(wh_data.values())
            n = noms.get(nid, {})
            lines.append(f"  {n.get('name', nid)}: {fmt(total)} {n.get('unit', 'шт')}")
        parts.append("ОСТАТКИ НА СКЛАДЕ:\n" + "\n".join(lines))

    debts = calc_debts(db)
    if debts:
        lines = []
        for c_id, amount in debts.items():
            name = get_contractor_name(db, c_id)
            if amount > 0:
                lines.append(f"  {name}: должен нам {fmt(amount)} сом")
            elif amount < 0:
                lines.append(f"  {name}: мы должны {fmt(-amount)} сом")
        if lines:
            parts.append("ДОЛГИ/ВЗАИМОРАСЧЁТЫ:\n" + "\n".join(lines))

    nom_list = db.get("trade", {}).get("nomenclature", [])
    if nom_list:
        lines = [f"  {n['name']}: цена {fmt(n.get('price', 0))} сом, себест. {fmt(n.get('cost', 0))} сом/{n.get('unit', 'шт')}" for n in nom_list]
        parts.append("НОМЕНКЛАТУРА И ЦЕНЫ:\n" + "\n".join(lines))

    contractors = db.get("trade", {}).get("contractors", [])
    if contractors:
        lines = [f"  {c['name']} (ИНН: {c.get('inn', '-')}, {c.get('kind', '')}, тел: {c.get('phone', '-')})" for c in contractors]
        parts.append("КОНТРАГЕНТЫ:\n" + "\n".join(lines))

    employees = db.get("trade", {}).get("employees", [])
    positions = {p["id"]: p["name"] for p in db.get("trade", {}).get("positions", [])}
    if employees:
        lines = [f"  {e['name']} — {positions.get(e.get('position', ''), '?')}, оклад {fmt(e.get('salary', 0))} сом" for e in employees]
        parts.append("СОТРУДНИКИ:\n" + "\n".join(lines))

    return "\n\n".join(parts)


AI_SYSTEM_PROMPT = """Ты — бухгалтерский AI-ассистент компании. Отвечаешь на вопросы по данным учёта.

ПРАВИЛА:
- Отвечай ТОЛЬКО на основе предоставленных данных. Не выдумывай цифры.
- Если в данных нет ответа, честно скажи что данных нет.
- Отвечай кратко, по делу, с конкретными цифрами.
- Используй эмодзи для структуры (💰📦📊🤝👥 и т.д.)
- НЕ используй Markdown форматирование (без **, без __, без ```). Только простой текст.
- Числа форматируй с пробелами (1 000 000).
- Валюта — сом.
- Язык — русский."""


def ask_ai(question: str, db: dict) -> str | None:
    if not ANTHROPIC_API_KEY or not HAS_ANTHROPIC:
        return None
    try:
        client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
        context = _prepare_data_context(db)
        response = client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=1024,
            system=AI_SYSTEM_PROMPT,
            messages=[{
                "role": "user",
                "content": f"ДАННЫЕ УЧЁТА:\n{context}\n\nВОПРОС: {question}"
            }],
        )
        text = response.content[0].text.strip()
        return _esc(text) if text else None
    except Exception as e:
        log.error(f"AI ошибка: {e}")
        return None


# ══════════════════════════════════════════════
#  Маршрутизация
# ══════════════════════════════════════════════

HANDLERS = [
    handle_summary, handle_period_report, handle_money, handle_debts, handle_price,
    handle_payroll, handle_org, handle_contracts, handle_warehouses,
    handle_stock, handle_goods, handle_contractors, handle_employees,
]


def process_question(text: str) -> str:
    text_lower = text.lower().strip()
    try:
        db = fetch_db()
    except Exception as e:
        log.error(f"Ошибка загрузки данных: {e}")
        return "⚠️ Не удалось загрузить данные\\. Попробуйте позже\\."

    for handler in HANDLERS:
        result = handler(db, text_lower)
        if result is not None:
            return result

    ai_answer = ask_ai(text, db)
    if ai_answer:
        return f"🤖 {ai_answer}"

    return None


# ══════════════════════════════════════════════
#  Telegram: кнопки быстрых действий
# ══════════════════════════════════════════════

MAIN_MENU = InlineKeyboardMarkup([
    [InlineKeyboardButton("📊 Сводка", callback_data="q:сводка"),
     InlineKeyboardButton("💰 Деньги", callback_data="q:сколько денег на счетах")],
    [InlineKeyboardButton("📦 Склад", callback_data="q:остатки на складе"),
     InlineKeyboardButton("📊 Долги", callback_data="q:кто нам должен")],
    [InlineKeyboardButton("📋 Товары", callback_data="q:список товаров"),
     InlineKeyboardButton("🤝 Контрагенты", callback_data="q:контрагенты")],
    [InlineKeyboardButton("👥 Сотрудники", callback_data="q:сотрудники"),
     InlineKeyboardButton("🏷 Прайс", callback_data="q:прайс-лист")],
    [InlineKeyboardButton("📈 График балансов", callback_data="chart:balances"),
     InlineKeyboardButton("📊 График склада", callback_data="chart:stock")],
    [InlineKeyboardButton("🔄 Приходы/Расходы", callback_data="chart:income"),
     InlineKeyboardButton("🥧 Долги (график)", callback_data="chart:debts")],
    [InlineKeyboardButton("📥 Excel: Балансы", callback_data="xlsx:balances"),
     InlineKeyboardButton("📥 Excel: Склад", callback_data="xlsx:stock")],
    [InlineKeyboardButton("📥 Excel: Долги", callback_data="xlsx:debts")],
])

HELP_MD = (
    f"Я могу ответить на:\n\n"
    f"  📦 _Сколько \\[товар\\] на складе?_\n"
    f"  💰 _Сколько денег на счетах?_\n"
    f"  🏷 _Сколько стоит \\[товар\\]?_\n"
    f"  📊 _Кто нам должен? / Наши долги_\n"
    f"  📈 _Сводка / Отчёт_\n"
    f"  📅 _Продажи за сентябрь / Расходы за неделю_\n"
    f"  📋 _Список товаров_\n"
    f"  🤝 _Контрагенты_\n"
    f"  👥 _Сотрудники_\n"
    f"  📝 _Договоры_\n"
    f"  🏢 _Реквизиты организации_\n"
    f"  💼 _Зарплатные ставки_\n\n"
    f"📊 _График баланса / График склада_\n"
    f"📥 _Выгрузи склад / Выгрузи балансы_\n\n"
    f"Или нажмите кнопку ниже 👇"
)


# ══════════════════════════════════════════════
#  Telegram: обработчики
# ══════════════════════════════════════════════

async def cmd_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    # Запоминаем chat_id для уведомлений
    chat_id = update.effective_chat.id
    if chat_id not in _notify_chat_ids:
        _notify_chat_ids.add(chat_id)
        log.info(f"Зарегистрирован chat_id для уведомлений: {chat_id}")
    await update.message.reply_text(
        f"Здравствуйте, {_esc(user.first_name)}\\! 👋\n\n"
        f"Я — бот бухгалтерии\\. Задайте вопрос или выберите раздел:\n\n"
        + HELP_MD,
        parse_mode=ParseMode.MARKDOWN_V2,
        reply_markup=MAIN_MENU,
    )

async def cmd_help(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(HELP_MD, parse_mode=ParseMode.MARKDOWN_V2, reply_markup=MAIN_MENU)

async def cmd_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("Выберите раздел:", reply_markup=MAIN_MENU)


async def _send_chart(chat, chart_bytes: bytes | None, caption: str, no_data_msg: str):
    if chart_bytes:
        await chat.send_photo(photo=chart_bytes, caption=caption)
    else:
        await chat.send_message(no_data_msg)


async def _send_xlsx(chat, xlsx_bytes: bytes | None, filename: str, no_data_msg: str):
    if xlsx_bytes:
        await chat.send_document(document=xlsx_bytes, filename=filename)
    else:
        await chat.send_message(no_data_msg)


async def handle_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = query.data
    chat = query.message.chat

    if data.startswith("q:"):
        question = data[2:]
        await chat.send_action(ChatAction.TYPING)
        answer = process_question(question)
        if not answer:
            answer = "Нет данных по этому запросу\\."
        if len(answer) > 4000:
            answer = answer[:4000] + "\n\n\\.\\.\\. \\(обрезано\\)"
        try:
            await query.message.reply_text(answer, parse_mode=ParseMode.MARKDOWN_V2, reply_markup=MAIN_MENU)
        except Exception:
            await query.message.reply_text(answer.replace("\\", ""), reply_markup=MAIN_MENU)

    elif data.startswith("chart:"):
        await chat.send_action(ChatAction.UPLOAD_PHOTO)
        try:
            db = fetch_db()
        except Exception:
            await chat.send_message("⚠️ Не удалось загрузить данные.")
            return
        chart_type = data[6:]
        if chart_type == "balances":
            await _send_chart(chat, chart_balances(db), "💰 Остатки по счетам", "Нет данных по счетам.")
        elif chart_type == "stock":
            await _send_chart(chat, chart_stock(db), "📦 Остатки на складе", "Нет данных по складу.")
        elif chart_type == "debts":
            await _send_chart(chat, chart_debts(db), "📊 Взаиморасчёты", "Долгов нет.")
        elif chart_type == "income":
            await _send_chart(chat, chart_income_expense(db), "🔄 Приходы и расходы", "Нет денежных операций.")

    elif data.startswith("xlsx:"):
        await chat.send_action(ChatAction.UPLOAD_DOCUMENT)
        try:
            db = fetch_db()
        except Exception:
            await chat.send_message("⚠️ Не удалось загрузить данные.")
            return
        xlsx_type = data[5:]
        today = datetime.now().strftime("%Y-%m-%d")
        if xlsx_type == "balances":
            await _send_xlsx(chat, export_balances_xlsx(db), f"Балансы_{today}.xlsx", "Нет данных по счетам.")
        elif xlsx_type == "stock":
            await _send_xlsx(chat, export_stock_xlsx(db), f"Склад_{today}.xlsx", "Нет данных по складу.")
        elif xlsx_type == "debts":
            await _send_xlsx(chat, export_debts_xlsx(db), f"Взаиморасчёты_{today}.xlsx", "Долгов нет.")


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text
    if not text:
        return
    log.info(f"[{update.effective_user.first_name}] {text}")

    await update.message.chat.send_action(ChatAction.TYPING)
    text_lower = text.lower().strip()

    # Запрос графика текстом
    chart_kw = {"график баланс": "balances", "график счет": "balances", "график счёт": "balances",
                "график склад": "stock", "график остат": "stock",
                "график долг": "debts", "график взаиморасч": "debts",
                "график приход": "income", "график расход": "income", "приходы расходы график": "income"}
    for kw, chart_type in chart_kw.items():
        if kw in text_lower:
            await update.message.chat.send_action(ChatAction.UPLOAD_PHOTO)
            try:
                db = fetch_db()
                funcs = {"balances": chart_balances, "stock": chart_stock, "debts": chart_debts, "income": chart_income_expense}
                captions = {"balances": "💰 Остатки по счетам", "stock": "📦 Остатки на складе", "debts": "📊 Взаиморасчёты", "income": "🔄 Приходы и расходы"}
                chart_bytes = funcs[chart_type](db)
                await _send_chart(update.message.chat, chart_bytes, captions[chart_type], "Нет данных для графика.")
            except Exception as e:
                log.error(f"Chart error: {e}")
                await update.message.reply_text("⚠️ Ошибка при создании графика.")
            return

    # Запрос Excel текстом
    xlsx_kw = {"выгрузи баланс": "balances", "выгрузи счет": "balances", "excel баланс": "balances",
               "выгрузи склад": "stock", "excel склад": "stock", "выгрузи остат": "stock",
               "выгрузи долг": "debts", "excel долг": "debts", "выгрузи взаиморасч": "debts",
               "выгрузи всё": "all", "excel всё": "all", "выгрузи все": "all", "excel все": "all"}
    for kw, xlsx_type in xlsx_kw.items():
        if kw in text_lower:
            await update.message.chat.send_action(ChatAction.UPLOAD_DOCUMENT)
            try:
                db = fetch_db()
                today = datetime.now().strftime("%Y-%m-%d")
                if xlsx_type == "all":
                    for name, func, fname in [("balances", export_balances_xlsx, f"Балансы_{today}.xlsx"),
                                               ("stock", export_stock_xlsx, f"Склад_{today}.xlsx"),
                                               ("debts", export_debts_xlsx, f"Взаиморасчёты_{today}.xlsx")]:
                        data = func(db)
                        if data:
                            await update.message.chat.send_document(document=data, filename=fname)
                    await update.message.reply_text("📥 Все отчёты выгружены\\!", parse_mode=ParseMode.MARKDOWN_V2)
                else:
                    funcs = {"balances": (export_balances_xlsx, f"Балансы_{today}.xlsx"),
                             "stock": (export_stock_xlsx, f"Склад_{today}.xlsx"),
                             "debts": (export_debts_xlsx, f"Взаиморасчёты_{today}.xlsx")}
                    func, fname = funcs[xlsx_type]
                    await _send_xlsx(update.message.chat, func(db), fname, "Нет данных для выгрузки.")
            except Exception as e:
                log.error(f"XLSX error: {e}")
                await update.message.reply_text("⚠️ Ошибка при создании файла.")
            return

    # Запрос Excel за период текстом
    if "выгрузи" in text_lower or "excel" in text_lower:
        period = parse_period(text_lower)
        if period:
            await update.message.chat.send_action(ChatAction.UPLOAD_DOCUMENT)
            try:
                db = fetch_db()
                date_from, date_to = period
                xlsx_data = export_period_xlsx(db, date_from, date_to)
                fname = f"Отчёт_{date_from.strftime('%d.%m')}-{date_to.strftime('%d.%m.%Y')}.xlsx"
                await _send_xlsx(update.message.chat, xlsx_data, fname, "Нет данных за этот период.")
            except Exception as e:
                log.error(f"Period XLSX error: {e}")
                await update.message.reply_text("⚠️ Ошибка при создании файла.")
            return

    # Обычный текстовый ответ
    answer = process_question(text)
    if not answer:
        answer = f"🤔 Не совсем понял вопрос\\.\n\n{HELP_MD}"

    if len(answer) > 4000:
        answer = answer[:4000] + "\n\n\\.\\.\\. \\(обрезано\\)"

    try:
        await update.message.reply_text(answer, parse_mode=ParseMode.MARKDOWN_V2, reply_markup=MAIN_MENU)
    except Exception as e:
        log.warning(f"Markdown error, fallback to plain: {e}")
        plain = re.sub(r'\\(.)', r'\1', answer)
        plain = re.sub(r'\*([^*]+)\*', r'\1', plain)
        await update.message.reply_text(plain, reply_markup=MAIN_MENU)

    log.info(f"[ответ] {answer[:80]}...")


# ══════════════════════════════════════════════
#  Уведомления
# ══════════════════════════════════════════════

_notify_chat_ids = set(NOTIFY_CHAT_IDS)
_last_doc_count = None
_last_low_balance_alert = 0

async def check_notifications(context: ContextTypes.DEFAULT_TYPE):
    """Проверяет данные и отправляет уведомления."""
    global _last_doc_count, _last_low_balance_alert
    if not _notify_chat_ids:
        return
    try:
        db = fetch_db()
    except Exception:
        return

    now = time.time()

    # 1. Новые документы
    bank_count = len(db.get("bankDocuments", []))
    cash_count = len(db.get("cashDocuments", []))
    trade_count = len(db.get("trade", {}).get("docs", []))
    total_docs = bank_count + cash_count + trade_count

    if _last_doc_count is not None and total_docs > _last_doc_count:
        diff = total_docs - _last_doc_count
        msg = f"🔔 Новых документов: {diff}\n\nВсего: банк {bank_count}, касса {cash_count}, торговля {trade_count}"
        for chat_id in _notify_chat_ids:
            try:
                await context.bot.send_message(chat_id=chat_id, text=msg)
            except Exception as e:
                log.warning(f"Notify error chat {chat_id}: {e}")
    _last_doc_count = total_docs

    # 2. Низкий баланс
    if now - _last_low_balance_alert > 3600:
        bal = calc_balances(db)
        low = []
        for acc_id, amount in bal.items():
            if 0 < amount < NOTIFY_MIN_BALANCE:
                low.append(f"  {get_account_name(db, acc_id)}: {fmt(amount)} сом")
        if low:
            msg = f"⚠️ Низкий баланс (менее {fmt(NOTIFY_MIN_BALANCE)} сом):\n\n" + "\n".join(low)
            for chat_id in _notify_chat_ids:
                try:
                    await context.bot.send_message(chat_id=chat_id, text=msg)
                except Exception as e:
                    log.warning(f"Notify error chat {chat_id}: {e}")
            _last_low_balance_alert = now


# ══════════════════════════════════════════════
#  Запуск
# ══════════════════════════════════════════════

async def post_init(app: Application):
    await app.bot.set_my_commands([
        BotCommand("start", "Начать работу"),
        BotCommand("menu", "Главное меню"),
        BotCommand("help", "Справка"),
    ])
    log.info("Команды бота зарегистрированы")


def main():
    if not TELEGRAM_TOKEN:
        log.error("TELEGRAM_TOKEN не задан!")
        return
    if not GITHUB_REPO:
        log.error("GITHUB_REPO не задан!")
        return

    log.info("Запуск бота...")
    app = Application.builder().token(TELEGRAM_TOKEN).post_init(post_init).build()

    app.add_handler(CommandHandler("start", cmd_start))
    app.add_handler(CommandHandler("help", cmd_help))
    app.add_handler(CommandHandler("menu", cmd_menu))
    app.add_handler(CallbackQueryHandler(handle_callback))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))

    # Уведомления каждые 5 минут
    app.job_queue.run_repeating(check_notifications, interval=300, first=30)

    log.info("Бот запущен!")
    app.run_polling(drop_pending_updates=True)


if __name__ == "__main__":
    main()
